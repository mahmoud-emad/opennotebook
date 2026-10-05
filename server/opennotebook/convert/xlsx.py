"""Excel (`.xlsx`) to Markdown, one GFM table per sheet, through openpyxl.

openpyxl already resolves shared strings, cell types and which number formats
are dates, so this module only decides how each value reads: whole numbers
without a trailing `.0`, fractions without binary noise, dates in ISO form. The
workbook is opened read-only and with cached values, so a formula reads as the
value Excel last computed, never as the formula. The first non-empty row is
taken as the header, since that is where a sheet's column names nearly always
are.
"""

from __future__ import annotations

import datetime as dt
import io
import math
import warnings
from collections.abc import Iterable, Sequence
from decimal import Decimal
from typing import Any

import openpyxl

from . import markdown
from .kinds import ConvertError, InputKind, Problem
from .ooxml import Package

KIND = InputKind.EXCEL

# Data rows written per sheet. A sheet past this is a dataset rather than a
# document; the rows written are enough to show what it holds, and the closing
# line says how much more there is.
MAX_ROWS = 2000


def convert(data: bytes) -> str:
    # openpyxl reads parts without a size limit, and parses worksheets with
    # the standard library's XML parser, so the whole archive is inflated once
    # under the budget first, refusing any part that declares a DTD. It costs a
    # second pass over the file, which is cheap next to holding a zip bomb in
    # memory.
    pkg = Package(data, KIND)
    try:
        pkg.check_budget(refuse_doctype=True)
    finally:
        pkg.close()

    with warnings.catch_warnings():
        # openpyxl warns about features it drops (data validation, unknown
        # extensions); none of them is text.
        warnings.simplefilter("ignore")
        try:
            book: Any = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
        except Exception as exc:
            raise ConvertError(KIND, Problem.NOT_VALID) from exc
        try:
            sections: list[str] = []
            for sheet in book.worksheets:
                try:
                    # Producers do not all write a correct used range, and
                    # read-only mode would trust it and cut the sheet short.
                    sheet.reset_dimensions()
                    rows = [[render(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                except Exception as exc:
                    raise ConvertError(KIND, Problem.DAMAGED) from exc
                tbl = sheet_table(rows)
                if tbl is not None:
                    sections.append(f"## {markdown.one_line(str(sheet.title))}\n\n{tbl}")
        finally:
            book.close()
    return "\n\n".join(sections)


def sheet_table(rows: Iterable[Sequence[str]]) -> str | None:
    """A sheet's rendered cells as a table, or `None` for a sheet with no
    values. Leading and trailing empty rows go, as do columns that are empty in
    every row at the right-hand edge; empty rows between data stay, because a
    gap in a sheet often separates one block of data from the next."""
    rows = [list(r) for r in rows]
    filled = [i for i, r in enumerate(rows) if any(r)]
    if not filled:
        return None
    rows = rows[filled[0] : filled[-1] + 1]

    width = max((max(i for i, c in enumerate(r) if c) + 1 for r in rows if any(r)), default=0)
    # A blank column on the left is trimmed as well, as Excel users often start
    # a table in column B.
    left = min((next(i for i, c in enumerate(r) if c) for r in rows if any(r)), default=0)
    rows = [r[:width][left:] for r in rows]

    omitted = max(len(rows) - 1 - MAX_ROWS, 0)
    out = markdown.table(rows[: MAX_ROWS + 1])
    if omitted > 0:
        noun = "row was" if omitted == 1 else "rows were"
        out += f"\n\n{omitted} more {noun} left out of this sheet."
    return out


def render(value: object) -> str:
    match value:
        case None:
            text = ""
        case bool():
            text = "TRUE" if value else "FALSE"
        case int():
            text = str(value)
        case float():
            text = number(value)
        case dt.datetime():
            text = date_time(value)
        case dt.date():
            text = value.isoformat()
        case dt.time():
            text = clock(value.hour * 3600 + value.minute * 60 + value.second, value.microsecond)
        case dt.timedelta():
            text = duration(value)
        case _:
            # Strings, and error values such as `#DIV/0!`, which openpyxl
            # also hands over as their text.
            text = str(value)
    return markdown.cell(text.replace("\r\n", "\n"))


def number(f: float) -> str:
    """A number as a person would type it. Excel stores every number as a
    double, so `0.1 + 0.2` is stored as `0.30000000000000004`; Excel itself
    shows at most 15 significant digits, and rounding to the same precision
    gives back the number the author saw."""
    if math.isnan(f):
        return "NaN"
    if math.isinf(f):
        return "inf" if f > 0 else "-inf"
    if f.is_integer() and abs(f) < 1e15:
        return str(int(f))
    rounded = float(f"{f:.14e}")
    # Positional, never scientific, as a person writes it.
    text = format(Decimal(repr(rounded)), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text == "-0" else text


def clock(seconds: int, microseconds: int = 0) -> str:
    total = seconds + round(microseconds / 1_000_000)
    return f"{total // 3600:02}:{total // 60 % 60:02}:{total % 60:02}"


def duration(d: dt.timedelta) -> str:
    """A duration (`[h]:mm`) as hours, minutes and seconds, since it is not a
    point in time."""
    total = round(d.total_seconds())
    return ("-" if total < 0 else "") + clock(abs(total))


def date_time(d: dt.datetime) -> str:
    """A date, or a date and time, in ISO 8601, with only as much of the time
    as it holds."""
    if (d.hour, d.minute, d.second, d.microsecond) == (0, 0, 0, 0):
        return f"{d.year:04}-{d.month:02}-{d.day:02}"
    if d.second == 0 and d.microsecond // 1000 == 0:
        return f"{d.year:04}-{d.month:02}-{d.day:02} {d.hour:02}:{d.minute:02}"
    return f"{d.year:04}-{d.month:02}-{d.day:02} {d.hour:02}:{d.minute:02}:{d.second:02}"
